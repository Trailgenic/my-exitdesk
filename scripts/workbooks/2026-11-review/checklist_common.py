import openpyxl, copy
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import Rule, DataBarRule
from openpyxl.styles.differential import DifferentialStyle
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.properties import PageSetupProperties

INK='FF1A1A18'; MUTED='FF6E6B64'; CREAM='FFF3E8D0'; PAPER='FFF7F5F0'; INPUT='FFFFF8E1'
F=lambda **k: Font(name='Arial', size=k.pop('sz',10), color=k.pop('color',INK), **k)
HDR_FILL=PatternFill('solid', fgColor=INK); CREAM_FILL=PatternFill('solid', fgColor=CREAM)
PAPER_FILL=PatternFill('solid', fgColor=PAPER); INPUT_FILL=PatternFill('solid', fgColor=INPUT)
WRAP=Alignment(wrap_text=True, vertical='top')
thin=Side(style='thin', color='FFD9D4C7'); BOX=Border(bottom=thin)

def input_fill(ws, cols, r1, r2):
    for col in cols:
        for r in range(r1, r2+1):
            ws[f'{col}{r}'].fill = INPUT_FILL

def dates(ws, cols, r1, r2):
    for col in cols:
        for r in range(r1, r2+1):
            ws[f'{col}{r}'].number_format = 'mmm d, yyyy'
        dv = DataValidation(type='date', operator='greaterThan', formula1='36526', allow_blank=True,
                            showErrorMessage=True, errorTitle='Date needed', error='Enter a date, for example 3/15/2027.')
        dv.add(f'{col}{r1}:{col}{r2}'); ws.add_data_validation(dv)

def dropdown(ws, rng, items, prompt=None):
    # remove any existing validation on exactly this range
    ws.data_validations.dataValidation = [d for d in ws.data_validations.dataValidation if str(d.sqref) != rng]
    dv = DataValidation(type='list', formula1='"' + ','.join(items) + '"', allow_blank=True, showErrorMessage=True,
                        errorTitle='Pick from the list', error='Choose one of: ' + ', '.join(items))
    if prompt: dv.promptTitle=None; dv.prompt=prompt; dv.showInputMessage=True
    dv.add(rng); ws.add_data_validation(dv)

def rebuild_text_cf(ws):
    """Existing containsText rules carry no formula, which Excel ignores. Rebuild them with formulas."""
    old = []
    for cf in ws.conditional_formatting:
        for r in cf.rules:
            old.append((str(cf.sqref), r))
    ws.conditional_formatting = type(ws.conditional_formatting)()
    for sqref, r in old:
        first = sqref.split(':')[0]
        if r.type == 'containsText' and r.text:
            nr = Rule(type='containsText', operator='containsText', text=r.text, dxf=r.dxf,
                      formula=[f'NOT(ISERROR(SEARCH("{r.text}",{first})))'])
            ws.conditional_formatting.add(sqref, nr)
        else:
            ws.conditional_formatting.add(sqref, r)

def text_cf(ws, rng, word, color, bold=True):
    first = rng.split(':')[0]
    ws.conditional_formatting.add(rng, Rule(type='containsText', operator='containsText', text=word,
        dxf=DifferentialStyle(font=Font(color=color, bold=bold)), formula=[f'NOT(ISERROR(SEARCH("{word}",{first})))']))

def print_setup(ws, landscape=True, title_rows=None, area=None, fit_height=0):
    ws.page_setup.orientation = 'landscape' if landscape else 'portrait'
    ws.page_setup.paperSize = ws.PAPERSIZE_LETTER
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = fit_height
    ws.page_margins.left = ws.page_margins.right = 0.4; ws.page_margins.top = ws.page_margins.bottom = 0.5
    if title_rows: ws.print_title_rows = title_rows
    if area: ws.print_area = area
    ws.oddFooter.center.text = '&A · page &P of &N'; ws.oddFooter.center.size = 8

def rename_header(ws, table, col_letter, new):
    old = ws[f'{col_letter}5'].value
    ws[f'{col_letter}5'].value = new
    for tc in ws.tables[table].tableColumns:
        if tc.name == old: tc.name = new
    return old

def clear_dashboard(ws, from_row):
    for mr in list(ws.merged_cells.ranges):
        if mr.min_row >= from_row: ws.unmerge_cells(str(mr))
    for row in ws.iter_rows(min_row=from_row, max_row=max(ws.max_row, from_row)):
        for c in row:
            c.value = None; c.font = F(); c.fill = PatternFill(); c.border = Border(); c.alignment = Alignment()
    ws.conditional_formatting = type(ws.conditional_formatting)()

class Dash:
    """Writes dashboard blocks top to bottom."""
    def __init__(self, ws, row):
        self.ws, self.r = ws, row
    def note(self, label, text):
        ws=self.ws; ws[f'A{self.r}']=label; ws[f'A{self.r}'].font=F(b=True); ws[f'A{self.r}'].fill=CREAM_FILL
        ws[f'A{self.r}'].alignment=WRAP
        ws[f'B{self.r}']=text; ws[f'B{self.r}'].font=F(); ws[f'B{self.r}'].alignment=WRAP
        ws.merge_cells(f'B{self.r}:F{self.r}')
        ws.row_dimensions[self.r].height = max(15, 15*(len(text)//120+1))
        self.r+=1
    def header(self, *cols):
        for i,t in enumerate(cols):
            c=self.ws.cell(self.r, i+1, t); c.font=F(b=True, color='FFFFFFFF'); c.fill=HDR_FILL; c.alignment=WRAP
        self.r+=1
    def metric(self, label, formula, use, fmt=None):
        ws=self.ws
        ws[f'A{self.r}']=label; ws[f'A{self.r}'].font=F(); ws[f'A{self.r}'].alignment=WRAP
        ws[f'B{self.r}']=formula; ws[f'B{self.r}'].font=F(sz=14, b=True); ws[f'B{self.r}'].fill=PAPER_FILL
        ws[f'B{self.r}'].alignment=Alignment(horizontal='left', vertical='center', wrap_text=True)
        if fmt: ws[f'B{self.r}'].number_format=fmt
        ws[f'C{self.r}']=use; ws[f'C{self.r}'].font=F(color=MUTED); ws[f'C{self.r}'].alignment=WRAP
        ws.merge_cells(f'C{self.r}:F{self.r}')
        for col in 'ABCDEF': ws[f'{col}{self.r}'].border=BOX
        self.r+=1; return self.r-1
    def row(self, values, fmts=None, bold_first=False):
        for i,v in enumerate(values):
            c=self.ws.cell(self.r, i+1, v); c.font=F(b=(bold_first and i==0)); c.alignment=WRAP; c.border=BOX
            if fmts and fmts[i]: c.number_format=fmts[i]
        self.r+=1; return self.r-1
    def gap(self, n=1): self.r+=n

def pct_bar(ws, rng):
    ws.conditional_formatting.add(rng, DataBarRule(start_type='num', start_value=0, end_type='num', end_value=1, color='FF8FA67E', showValue=True))
