import openpyxl
R='/home/claude/my-exitdesk/public/resources/'; O='/tmp/build/tools/'
def ver(wb, old='1.0', new='1.0.1'):
    n=0
    for ws in wb:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value,str) and not c.value.startswith('=') and ('v'+old in c.value or 'Version '+old in c.value):
                    c.value=c.value.replace('v'+old,'v'+new).replace('Version '+old,'Version '+new); n+=1
    return n
# Capital allocation
wb=openpyxl.load_workbook(R+'MikeYe-Capital-Allocation-Deal-Affordability-v1.0.xlsx')
wb['Deal cash flow']['D39']='=IF(OR(COUNT(D30:I30)<>6,D30>=0,MAX(E30:I30)<=0,SUMPRODUCT(--((D30:H30<0)<>(E30:I30<0)))>1),"n.a.",IFERROR(IRR(D30:I30),"n.a."))'
wb['Dashboard']['D10']=("=IF(ISNUMBER('Capacity'!D37),'Capacity'!D37,IF('Capacity'!D38<>\"Capacity calculated\",'Capacity'!D38,"
  "IF('Financing'!I22<>\"Funding confirmed\",\"Open: funding not committed\",IF('Financing'!I7>'Capacity'!D35,\"Open: proposed debt above committed limit\","
  "IF('Financing'!I19>'Controls'!D31,\"Open: dilution above limit\",\"Open inputs\")))))")
g=wb['Guide']; g['D20']=g['D20'].value.replace('a nonnegative initial flow or negative future total flows','a nonnegative initial flow, no positive future flow, or cash flows that change sign more than once')
wb['Controls']['F10']='1, 2 or 3: the three funding options named across the top of the Financing tab. Proposed debt is distinct from committed debt capacity.'
print('cap ver',ver(wb)); wb.save(O+'MikeYe-Capital-Allocation-Deal-Affordability-v1.0.1.xlsx')
# Synergy
wb=openpyxl.load_workbook(R+'MikeYe-Synergy-Underwriting-Value-Bridge-v1.0.xlsx'); a=wb['Assumptions']
a['D9']=22000; a['D12']='Assumption: CEO and board'; a['D13']=21000; a['D14']=20500
for c in ('D9','D13','D14'): a[c].number_format=a['D7'].number_format; a[c].font=a['D7'].font.copy(); a[c].fill=a['D7'].fill.copy()
a['D12'].font=a['D7'].font.copy(); a['D12'].fill=a['D7'].fill.copy()
a['E10']='Enter 1 if the buyer can use the tax deductions now, 0 if not'
wb['Initiatives']['C3']=('First pass: replace the three Assumption examples. Use all 10 fixed IDs and enter an explicit zero where appropriate. '
  'Dis-synergies count in full from close, even if the initiative starts later. That is deliberate.')
wb['Guide']['D16']=('Blank or invalid economic inputs hold back the cash forecast and price outputs. An explicit zero is accepted. '
  'Open findings stay visible after numbers are entered. When the operating owner has resolved a finding, type None in Missing input / finding.')
print('syn ver',ver(wb)); wb.save(O+'MikeYe-Synergy-Underwriting-Value-Bridge-v1.0.1.xlsx')
# LOI
wb=openpyxl.load_workbook(R+'MikeYe-LOI-Economics-Risk-Allocator-v1.0.xlsx'); d=wb['Dashboard']
d['F33']='Keep the private ceiling fixed as terms move.'
d['C11']='Buyer cost of equity consideration (today\'s dollars)'
d['C12']='Buyer cost as enterprise value (today\'s dollars)'
d['C13']='Enterprise value if every earnout pays in full'
wb['Assumptions']['F20']="Compare with 'Enterprise value if every earnout pays in full' on the Dashboard."
print('loi ver',ver(wb)); wb.save(O+'MikeYe-LOI-Economics-Risk-Allocator-v1.0.1.xlsx')
# Mandate
wb=openpyxl.load_workbook(R+'MikeYe-Acquisition-Mandate-Target-Screen-v1.0.xlsx')
calc=openpyxl.load_workbook('/tmp/recalc/out/MikeYe-Acquisition-Mandate-Target-Screen-v1.0.xlsx',data_only=True)['Screen']
s=wb['Screen']
for r in range(31,79):
    s[f'O{r}']=1; s[f'P{r}']=calc[f'Q{r}'].value
    s[f'O{r}']._style=s['O7']._style; s[f'P{r}']._style=s['P7']._style
t=wb['Targets']
for r in range(11,19): t[f'AH{r}']=1; t[f'AH{r}']._style=t['AH7']._style
g=wb['Guide']
g['D12']=(g['D12'].value+' The workbook uses all six criteria rows. If you need fewer, keep the extra rows filled in and set Required? to No.')
g['D27']=('Fill an unused target slot, assign a unique priority and complete its six Screen rows. Reviewed revision and basis (columns O and P) are pre-filled for the current mandate; '
          'update them only after the mandate changes. The dashboard updates automatically. A missing required pair stays open; duplicate pairs are invalid.')
g['D30']='Use this workbook to set the mandate and the shortlist. Then use the diligence, valuation, synergy and LOI tools on mikeye.com as each decision comes up.'
g['C30']='What comes next'
wb['Dashboard'].freeze_panes='A5'
print('man ver',ver(wb)); wb.save(O+'MikeYe-Acquisition-Mandate-Target-Screen-v1.0.1.xlsx')
