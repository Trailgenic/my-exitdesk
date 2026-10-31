import sys; sys.path.insert(0,'/tmp/build')
from checklist_common import *
from openpyxl.worksheet.table import TableColumn
SRC='/home/claude/my-exitdesk/public/resources/MikeYe-Post-Deal-Integration-Checklist-v1.0.xlsx'
OUT=sys.argv[1]
wb=openpyxl.load_workbook(SRC)
HML=['High','Medium','Low']
ip=wb['Integration Plan']; R1,R2=6,109
dropdown(ip,f'L{R1}:L{R2}',HML)
input_fill(ip,['H','I','J','N','O','P','Q','R'],R1,R2); dates(ip,['I'],R1,R2)
ip['A3']='Protect revenue while integrating cost. Each row is one action: give it an owner and a date before changing anything that already works. Yellow columns are yours to fill in.'
rebuild_text_cf(ip); print_setup(ip,title_rows='5:5',area=f'A5:Q{R2}')

d1=wb['Day 1 Readiness']; input_fill(d1,['E','F','G','H','I'],6,17); print_setup(d1,title_rows='5:5')
text_cf(d1,'F6:F17','Blocked','FF8B2B24'); text_cf(d1,'F6:F17','Ready','FF315B37')
wd=wb['Workflow Decisions']; dropdown(wd,'E6:E19',HML); input_fill(wd,list('CDEFGHIJKLMNOPQ'),6,19); print_setup(wd,title_rows='5:5')

sy=wb['Synergies']
for r in range(6,46):
    for col in 'FGN':
        if sy[f'{col}{r}'].value in ('0',0): sy[f'{col}{r}'].value=None
rename_header(sy,'SynergyTable','E','Baseline (annual)')
rename_header(sy,'SynergyTable','F','Target (annual)')
rename_header(sy,'SynergyTable','N','Realized to date (annual run-rate)')
# add two calculated columns Q, R inside the table
src=sy['P5']
for col,name,w in [('Q','Annual benefit',16),('R','Realized %',12)]:
    c=sy[f'{col}5']; c.value=name; c._style=copy.copy(src._style)
    sy.column_dimensions[col].width=w
t=sy.tables['SynergyTable']; t.ref='A5:R45'
t.tableColumns.append(TableColumn(id=len(t.tableColumns)+1,name='Annual benefit'))
t.tableColumns.append(TableColumn(id=len(t.tableColumns)+1,name='Realized %'))
t.autoFilter.ref='A5:R45'
for r in range(6,46):
    sy[f'Q{r}']=f'=IF(OR(E{r}="",F{r}=""),"",IF(OR(B{r}="Revenue",B{r}="Dis-synergy"),F{r}-E{r},E{r}-F{r}))'
    sy[f'R{r}']=f'=IF(OR(Q{r}="",Q{r}=0,N{r}=""),"",N{r}/Q{r})'
    for col,fmt in [('E','#,##0;(#,##0)'),('F','#,##0;(#,##0)'),('G','#,##0;(#,##0)'),('N','#,##0;(#,##0)'),('Q','#,##0;(#,##0);-'),('R','0%')]:
        sy[f'{col}{r}'].number_format=fmt
    for col in 'QR': sy[f'{col}{r}']._style=copy.copy(sy[f'P{r}']._style); sy[f'{col}{r}'].number_format='#,##0;(#,##0);-' if col=='Q' else '0%'
input_fill(sy,list('BCDEFGHIJKLMNOP'),6,45); dates(sy,['H','I'],6,45)
sy['A3']=('Enter only underwritten synergies. Baseline and Target are annual amounts in the same units (for a cost line: today\'s cost and the cost after the change; '
          'for revenue: today\'s revenue and the target). Annual benefit and Realized % calculate themselves. Enter a dis-synergy as lost revenue.')
print_setup(sy,title_rows='5:5')

rk=wb['Risks & Decisions']; dropdown(rk,'G6:G45',HML); dropdown(rk,'H6:H45',['Investigate','Price','Protect','Walk','Approve','Escalate','Defer'])
input_fill(rk,list('BCDEFGHIJKLM'),6,45); dates(rk,['J'],6,45); print_setup(rk,title_rows='5:5')
text_cf(rk,'F6:F45','Critical','FF8B2B24')

ds=wb['Dashboard']; clear_dashboard(ds,5)
ds['A3']='Version 1.1 · November 2026 · 104 integration actions, 12 Day 1 gates, 14 workflow decisions, a synergy register and a risk log.'
d=Dash(ds,5)
d.note('How to use','Clear the Day 1 gates first. Then work the Integration Plan by phase: owner and date on every action, and no change to a working process until its evidence is in. Set Reviewer decision to Remove for actions that don\'t apply; they drop out of the counts.')
d.note('Rule','Protect revenue while integrating cost. Count a saving only once the old cost has stopped and the new running cost is in.')
d.gap()
R='ReadinessTable'; Rlive=f'{R}[Reviewer decision],"<>Remove",{R}[Status],"<>N/A"'
d.header('Day 1','Now','What it tells you')
d.metric('Day 1 verdict',f'=IF(COUNTIFS({Rlive})=COUNTIFS({R}[Status],"Ready",{R}[Reviewer decision],"<>Remove"),"Day 1 ready","Not ready: "&(COUNTIFS({Rlive})-COUNTIFS({R}[Status],"Ready",{R}[Reviewer decision],"<>Remove"))&IF(COUNTIFS({Rlive})-COUNTIFS({R}[Status],"Ready",{R}[Reviewer decision],"<>Remove")=1," gate open"," gates open"))','Every in-scope gate on the Day 1 tab must read Ready.')
d.metric('Gates ready',f'=COUNTIFS({R}[Status],"Ready",{R}[Reviewer decision],"<>Remove")&" of "&COUNTIFS({Rlive})','Legal control, cash, payroll, customers, systems and the rest.')
d.metric('Gates blocked',f'=COUNTIFS({R}[Status],"Blocked",{R}[Reviewer decision],"<>Remove")','Fix or escalate before close.')
d.gap()
T='IntegrationPlanTable'; live=f'{T}[Reviewer decision],"<>Remove"'
d.header('Integration plan','Now','What it tells you')
rs=d.metric('Actions in scope',f'=COUNTIFS({live},{T}[Status],"<>N/A")','Actions still in play.')
rc=d.metric('Complete',f'=COUNTIFS({T}[Status],"Complete",{live})','Done, with the evidence in the Ready column.')
d.metric('% complete',f'=IFERROR(B{rc}/B{rs},0)','Share of in-scope actions done.','0%')
d.metric('Blocked',f'=COUNTIFS({T}[Status],"Blocked",{live})','Stuck. Name the fix or escalate.')
d.metric('Overdue',f'=COUNTIFS({T}[Target date],"<"&TODAY(),{T}[Status],"<>Complete",{T}[Status],"<>N/A",{live})','Past the target date and not done.')
d.metric('Open Critical actions with no owner',f'=COUNTIFS({T}[Risk],"Critical",{T}[Status],"<>Complete",{T}[Status],"<>N/A",{T}[Accountable owner],"",{live})','Assign these first.')
d.metric('Open High revenue-impact actions',f'=COUNTIFS({T}[Revenue impact],"High",{T}[Status],"<>Complete",{T}[Status],"<>N/A",{live})','Where customers or revenue could be hurt. Go slow here.')
d.gap()
S='SynergyTable'; Slive=f'{S}[Reviewer decision],"<>Remove",{S}[Status],"<>Removed"'
d.header('Synergies','Amount','What it tells you')
ru=d.metric('Underwritten annual benefit',f'=SUMIFS({S}[Annual benefit],{Slive})','Net of dis-synergies, from the Synergies tab.','#,##0;(#,##0);-')
rr=d.metric('Realized run-rate to date',f'=SUMIFS({S}[Realized to date (annual run-rate)],{Slive})','Count only savings where the old cost has actually stopped.','#,##0;(#,##0);-')
d.metric('% realized',f'=IFERROR(B{rr}/B{ru},0)','Realized run-rate against what was underwritten.','0%')
d.metric('One-time cost to achieve',f'=SUMIFS({S}[One-time cost],{Slive})','Total cost to deliver the synergies.','#,##0;(#,##0);-')
d.gap()
K='IntegrationRiskTable'; Kopen=f'{K}[Status],"<>Resolved",{K}[Status],"<>Accepted",{K}[Issue / decision],"<>",{K}[Reviewer decision],"<>Remove"'
d.header('Risks and decisions','Open','What it tells you')
d.metric('Open issues',f'=COUNTIFS({Kopen})','Logged on the Risks & Decisions tab.')
d.metric('Critical severity',f'=COUNTIFS({K}[Severity],"Critical",{Kopen})','Open issues rated Critical.')
d.metric('Escalated',f'=COUNTIFS({K}[Status],"Escalated",{K}[Reviewer decision],"<>Remove")','Waiting on a decision above the integration lead.')
d.metric('Overdue',f'=COUNTIFS({K}[Due date],"<"&TODAY(),{Kopen})','Past due and still open.')
d.gap()
d.header('Progress by phase','Actions','Complete','Blocked','% complete')
first=d.r
for p in ['Diligence','Sign-to-Close','Day 1','Days 2-30','Days 31-60','Days 61-100','Post-100']:
    r=d.r
    d.row([p,f'=COUNTIFS({T}[Phase],A{r},{live})',f'=COUNTIFS({T}[Phase],A{r},{T}[Status],"Complete",{live})',
           f'=COUNTIFS({T}[Phase],A{r},{T}[Status],"Blocked",{live})',f'=IFERROR(C{r}/(B{r}-COUNTIFS({T}[Phase],A{r},{T}[Status],"N/A",{live})),0)'],fmts=[None,'0','0','0','0%'])
pct_bar(ds,f'E{first}:E{d.r-1}')
d.gap()
d.header('Disposition','Meaning')
for k,t in [('Preserve','Keep a proven workflow running as-is while the buyer learns.'),('Protect','Add redundancy, rights, controls or retention around fragile value.'),
            ('Combine','Join complementary capabilities only when the value and dependencies are clear.'),('Standardize','Move to a common process where the evidence supports it.'),
            ('Retire','Remove duplication only after continuity, migration and rollback are proven.')]:
    r=d.r; ds[f'A{r}']=k; ds[f'A{r}'].font=F(b=True); ds[f'B{r}']=t; ds[f'B{r}'].font=F(); ds[f'B{r}'].alignment=WRAP; ds.merge_cells(f'B{r}:F{r}'); d.r+=1
for col,w in zip('ABCDEF',[34,18,16,16,16,30]): ds.column_dimensions[col].width=w
print_setup(ds,landscape=False,area=f'A2:F{d.r-1}',fit_height=1)
wb.properties.title='Post-Deal Integration Checklist v1.1'; wb.properties.creator='Mike Ye'
wb.active=0; wb.save(OUT); print('saved')
