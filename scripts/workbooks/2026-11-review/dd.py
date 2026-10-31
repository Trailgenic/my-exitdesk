import sys; sys.path.insert(0,'/tmp/build')
from checklist_common import *
from openpyxl.comments import Comment
SRC='/home/claude/my-exitdesk/public/resources/MikeYe-MA-Due-Diligence-Checklist-v1.0.xlsx'
OUT=sys.argv[1]
wb=openpyxl.load_workbook(SRC)
IMP=['Investigate','Price','Protect','Walk']
ck=wb['Checklist']; R1,R2=6,117
dropdown(ck,f'J{R1}:J{R2}',IMP)
rename_header(ck,'DDChecklistTable','O','Risk found')
ck['O5'].comment=Comment('Fill this in after the work. Priority is how much the item matters before you look. Risk found is what you actually found.','Mike Ye')
input_fill(ck,['L','M','N','O','P','Q','R'],R1,R2)
ck['A3']='Filter by phase and workstream. Each row is a question to answer with evidence. Yellow columns are yours to fill in; send anything material to the Issue Log.'
rebuild_text_cf(ck); print_setup(ck,title_rows='5:5',area=f'A5:Q{R2}')

il=wb['Issue Log']
dropdown(il,'H6:H45',IMP); dropdown(il,'G6:G45',['High','Medium','Low'])
input_fill(il,list('BCDEFGHIJKLMNO'),6,45); dates(il,['M'],6,45)
text_cf(il,'H6:H45','Walk','FF8B2B24'); text_cf(il,'F6:F45','Critical','FF8B2B24')
print_setup(il,title_rows='5:5')

wm=wb['Workflow Map']
dropdown(wm,'O6:O19',['Transfers','Transfers with action','At risk','Does not transfer'])
dropdown(wm,'P6:P19',['High','Partial','Low','Not permitted'])
dropdown(wm,'Q6:Q19',['Preserve','Protect','Combine','Standardize','Retire'])
input_fill(wm,list('CDEFGHIJKLMNOPQRS'),6,19); print_setup(wm,title_rows='5:5')
text_cf(wm,'O6:O19','Does not transfer','FF8B2B24'); text_cf(wm,'O6:O19','At risk','FF6D531D')

# ---- Follow One Sale ----
fs=wb.create_sheet('Follow One Sale',1)
fs.sheet_view.showGridLines=False
fs['A2']='Follow One Sale'; fs['A2'].font=F(sz=16,b=True)
fs['A3']=('Pick one real customer sale from the last 12 months and trace it end to end, from the agreement to the cash in the bank. '
          'Do it with the person who actually does the work, not the founder. Then repeat for a second sale from a different customer or channel.')
fs['A3'].font=F(color=MUTED); fs['A3'].alignment=WRAP; fs.merge_cells('A3:G3'); fs.row_dimensions[3].height=42
fs['A4']='Customer / sale traced:'; fs['A4'].font=F(b=True); fs['C4'].fill=INPUT_FILL; fs.merge_cells('C4:E4')
hdr=['Step','What should exist','Who did it','What you found','Matches?','Gap or question','Send to']
for i,h in enumerate(hdr):
    c=fs.cell(6,i+1,h); c.font=F(b=True,color='FFFFFFFF'); c.fill=HDR_FILL; c.alignment=WRAP
steps=[('1. Agreement','Signed contract, order form or accepted quote with price, term and renewal or cancellation terms.'),
 ('2. Order','The order as entered in the system, matching the agreed price and quantity.'),
 ('3. Delivery','Proof the product shipped or the service was performed, with a date.'),
 ('4. Invoice','Invoice that matches the order and the delivery date.'),
 ('5. Revenue entry','The sale in the general ledger, in the right month and the right account.'),
 ('6. Cash receipt','Bank statement line showing the cash arrived, and when.'),
 ('7. Renewal or repeat','Evidence this customer bought again or renewed, and who managed that.'),
 ('8. Exceptions','Credits, discounts, rebates or side letters that changed what the business kept.')]
for i,(s,w) in enumerate(steps):
    r=7+i
    fs.cell(r,1,s).font=F(b=True); fs.cell(r,2,w).font=F()
    for col in range(1,8):
        c=fs.cell(r,col); c.alignment=WRAP; c.border=BOX
        if col>=3: c.fill=INPUT_FILL
    fs.row_dimensions[r].height=44
dropdown(fs,'E7:E14',['Yes','Partly','No','Not seen'])
dropdown(fs,'G7:G14',['Nothing to send','Issue Log','Price','Agreement','Integration plan'])
text_cf(fs,'E7:E14','No','FF8B2B24'); text_cf(fs,'E7:E14','Not seen','FF8B2B24')
fs['A16']='What to look for'; fs['A16'].font=F(b=True); fs['A16'].fill=CREAM_FILL
tips=['A timing difference between invoice, revenue and cash is normal. Ask for the explanation before you call it a problem.',
      'If only the founder can answer a step, write that down. It is a transferability finding, not a paperwork gap.',
      'If a step has no document at all, mark it Not seen and log it. Missing evidence is a finding.']
for i,t in enumerate(tips):
    c=fs.cell(17+i,1,'• '+t); c.font=F(); c.alignment=WRAP; fs.merge_cells(start_row=17+i,start_column=1,end_row=17+i,end_column=7); fs.row_dimensions[17+i].height=30
for col,w in zip('ABCDEFG',[20,44,18,40,12,36,18]): fs.column_dimensions[col].width=w
fs.freeze_panes='A7'; print_setup(fs,area='A2:G19',fit_height=1)

# ---- Dashboard ----
ds=wb['Dashboard']; clear_dashboard(ds,5)
ds['A3']='Version 1.1 · November 2026 · 112 diligence items, a sale-tracing worksheet, a workflow-transferability map and an issue log.'
d=Dash(ds,5)
d.note('How to use','Start with Follow One Sale. Then work the Checklist by phase and workstream: assign owners, update Status and record what you found. Anything material goes to the Issue Log with a decision. Set Reviewer decision to Remove for items that don\'t apply; they drop out of the counts.')
d.note('Decision rule','Every material finding becomes Investigate, Price, Protect or Walk, and then carries into the price, the agreement or the integration plan.')
d.gap()
T='DDChecklistTable'; live=f'{T}[Reviewer decision],"<>Remove"'
d.header('Checklist progress','Now','What it tells you')
rs=d.metric('Items in scope',f'=COUNTIFS({live},{T}[Status],"<>N/A")','Items still in play.')
rc=d.metric('Complete',f'=COUNTIFS({T}[Status],"Complete",{live})','Answered, with a finding recorded.')
d.metric('% complete',f'=IFERROR(B{rc}/B{rs},0)','Share of in-scope items done.','0%')
d.metric('Blocked',f'=COUNTIFS({T}[Status],"Blocked",{live})','Usually waiting on the seller. Chase or escalate.')
d.metric('Open Critical items with no owner',f'=COUNTIFS({T}[Priority],"Critical",{T}[Status],"<>Complete",{T}[Status],"<>N/A",{T}[Owner],"",{live})','Assign these first.')
d.metric('Workflows mapped','=COUNTIFS(DDWorkflowTable[Change-of-control transferability],"<>")&" of "&COUNTIFS(DDWorkflowTable[Reviewer decision],"<>Remove")','Workflows with a transferability call on the Workflow Map.')
d.gap()
I='DDIssueTable'; openI=f'{I}[Status],"<>Resolved",{I}[Status],"<>Accepted",{I}[Finding],"<>",{I}[Reviewer decision],"<>Remove"'
d.header('Open issues','Open','What it means')
defs=[('Walk','The downside, a transfer failure or an integrity concern breaks the thesis.'),
      ('Price','Changes standalone value, required investment or financing.'),
      ('Protect','Handle through structure, contract, consent, retention, control or integration.'),
      ('Investigate','Evidence is missing or the claim still needs testing. Unknowns stay open until evidence closes them.')]
for k,t in defs:
    rr=d.metric(k,f'=COUNTIFS({I}[Decision impact],A{d.r},{openI})',t); ds[f'A{rr}'].font=F(b=True)
d.metric('Critical severity',f'=COUNTIFS({I}[Severity],"Critical",{openI})','Open issues rated Critical.')
d.metric('Overdue',f'=COUNTIFS({I}[Due date],"<"&TODAY(),{openI})','Past due and still open.')
d.gap()
d.header('Progress by phase','Items','Complete','Blocked','% complete')
first=d.r
for p in ['Pre-LOI','LOI','Confirmatory','Pre-close']:
    rr=d.r
    d.row([p,f'=COUNTIFS({T}[Phase],A{rr},{live})',f'=COUNTIFS({T}[Phase],A{rr},{T}[Status],"Complete",{live})',
           f'=COUNTIFS({T}[Phase],A{rr},{T}[Status],"Blocked",{live})',f'=IFERROR(C{rr}/(B{rr}-COUNTIFS({T}[Phase],A{rr},{T}[Status],"N/A",{live})),0)'],fmts=[None,'0','0','0','0%'])
pct_bar(ds,f'E{first}:E{d.r-1}')
for col,w in zip('ABCDEF',[34,16,16,16,16,30]): ds.column_dimensions[col].width=w
print_setup(ds,landscape=False,area=f'A2:F{d.r-1}',fit_height=1)
wb.properties.title='M&A Due Diligence Checklist v1.1'; wb.properties.creator='Mike Ye'
wb.active=0; wb.save(OUT); print('saved')
